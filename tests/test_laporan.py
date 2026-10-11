import io

import numpy as np
from openpyxl import load_workbook

import adimology as ad
import fundamental as fd
import laporan as lp
import teknikal as tk
from fakes import setup_case
from test_fundamental import _company
from test_teknikal import ohlc


def _inputs(liab_scale=0.5):
    r = ad.analyze(setup_case(), ad.Params(kode="TEST"))
    x = np.arange(320)
    rk = tk.rekomendasi(ohlc(np.linspace(80, 120, 320) + 6 * np.sin(x * 2 * np.pi / 30)), {"status": "AKUMULASI"})
    f = fd.analisa(harga=r.harga, **_company(liab_scale=liab_scale))
    return r, rk, f


def _nilai(kat):
    return {"kategori": kat, "skor": 4, "skor_maks": 7, "jebakan": ["contoh jebakan"] if kat.startswith("WASPADA") else [],
            "kriteria": [], "info": []}


def test_keputusan_rules():
    r, rk, f = _inputs()
    beli = dict(rk, verdict="BELI DI AREA ENTRY")
    assert lp.keputusan(r, beli, f, nilai=_nilai("AKUMULASI KUAT"))[0] == "BELI"
    kep, alasan = lp.keputusan(r, beli, f, nilai=_nilai("AKUMULASI LEMAH / BELUM SEARAH"))
    assert kep == "TUNGGU" and any("jejak akumulasi" in a for a in alasan)
    # jebakan ringan (harga dikejar) hanya menahan; jebakan penjualan/distribusi = veto
    assert lp.keputusan(r, beli, f, nilai=_nilai("WASPADA: TANDA DISTRIBUSI / JEBAKAN"))[0] == "TUNGGU"
    n_berat = dict(_nilai("AKUMULASI TERINDIKASI"), jebakan=["Bandar menjual 5.000 lot di pasar negosiasi"])
    kep_b, al_b = lp.keputusan(r, beli, f, nilai=n_berat)
    assert kep_b == "HINDARI / JUAL" and any("veto" in a for a in al_b)
    # bandar = sensor risiko + penguat: jejak akumulasi (terkonfirmasi/terindikasi/absorpsi) + teknikal di area beli -> BELI; tidak konklusif -> TUNGGU
    kep_k, al_k = lp.keputusan(r, beli, f, nilai=_nilai("AKUMULASI TERKONFIRMASI (perilaku harga)"))
    assert kep_k == "BELI" and any("TERKONFIRMASI" in a for a in al_k)
    kep_t, al_t = lp.keputusan(r, beli, f, nilai=_nilai("AKUMULASI TERINDIKASI"))
    assert kep_t == "BELI" and any("BELUM terkonfirmasi" in a for a in al_t)
    assert lp.keputusan(r, beli, f, nilai=_nilai("ABSORPSI SUPPLY"))[0] == "BELI"
    assert lp.keputusan(r, beli, f, nilai=_nilai("TIDAK KONKLUSIF"))[0] == "TUNGGU"
    # veto ledger: DISTRIBUTING/MARKDOWN dengan confidence cukup -> HINDARI; confidence rendah hanya menahan BELI
    sg = dict(episode_status="DISTRIBUTING", confidence_score=0.8, bandar_score=3.0)
    assert lp.keputusan(r, beli, f, nilai=_nilai("AKUMULASI TERINDIKASI"), ctx=dict(mode="lengkap", ledger=sg))[0] == "HINDARI / JUAL"
    assert lp.keputusan(r, beli, f, nilai=_nilai("AKUMULASI TERINDIKASI"), ctx=dict(mode="lengkap", ledger=dict(sg, confidence_score=0.4)))[0] == "TUNGGU"
    assert lp.keputusan(r, beli, f, nilai=_nilai("AKUMULASI TERINDIKASI"), ctx=dict(mode="lengkap", ledger=dict(sg, episode_status="ACCUMULATING", bandar_score=8)))[0] == "BELI"
    assert lp.keputusan(r, beli, f, nilai=_nilai("AKUMULASI TERINDIKASI"), ctx=dict(mode="cepat"))[0] == "TUNGGU"
    assert lp.keputusan(r, beli, f, nilai=_nilai("DISTRIBUSI TERINDIKASI"))[0] == "HINDARI / JUAL"
    assert lp.keputusan(r, dict(rk, verdict="HINDARI"), f, nilai=_nilai("AKUMULASI KUAT"))[0] == "HINDARI / JUAL"
    _, _, f_bad = _inputs(liab_scale=1.3)                         # utang > aset menang atas semuanya
    kep, alasan = lp.keputusan(r, beli, f_bad, nilai=_nilai("AKUMULASI KUAT"))
    assert kep == "HINDARI / JUAL" and any("LEBIH BESAR" in a for a in alasan)
    assert any("Skenario Adimology" in a for a in lp.keputusan(r, None, None)[1])


def test_pdf_and_excel_are_built():
    r, rk, f = _inputs()
    pdf = lp.pdf_bytes(r, rk, f, "2026-10-10 15:00 WIB")
    assert pdf[:4] == b"%PDF" and len(pdf) > 20_000             # termasuk grafik
    wb = load_workbook(io.BytesIO(lp.excel_bytes(r, rk, f, "2026-10-10 15:00 WIB")))
    assert {"Ringkasan", "Teknikal", "Broker", "PnL", "Neraca", "Arus kas", "Valuasi"} <= set(wb.sheetnames)
    ring = {row[0].value: row[1].value for row in wb["Ringkasan"].iter_rows(min_row=2)}
    assert ring["Kode"] == "TEST" and ring["Keputusan"] in lp.WARNA_KEP
    pnl = wb["PnL"]
    fills = {pnl.cell(row=i, column=j).fill.fgColor.rgb for i in range(2, pnl.max_row + 1) for j in range(2, pnl.max_column + 1)}
    assert any(str(c).endswith("DCFCE7") for c in fills)          # sel hijau ikut ke Excel


def test_export_with_bpjp_bsjp():
    import intraday as idt
    from test_intraday import C, bars
    from openpyxl import load_workbook
    r, rk, f = _inputs()
    o = idt.statistik(idt.get_bars(C(bars(61)), "TEST", None))
    assert lp.pdf_bytes(r, rk, f, "x", o)[:4] == b"%PDF"
    wb = load_workbook(io.BytesIO(lp.excel_bytes(r, rk, f, "x", o)))
    assert {"BSJP", "BPJP", "BSJP data harian"} <= set(wb.sheetnames)


def test_export_without_technical_or_fundamental():
    r, _, _ = _inputs()
    assert lp.pdf_bytes(r, None, None, "x")[:4] == b"%PDF"
    assert "Ringkasan" in load_workbook(io.BytesIO(lp.excel_bytes(r, None, None, "x"))).sheetnames


def test_beli_needs_hengky_core_rules():
    r, rk, f = _inputs()
    beli = dict(rk, verdict="BELI DI AREA ENTRY")
    gagal = {"checklist": {"lolos_inti": False, "hasil": "TAHAN: 2 dari 4 aturan inti belum terpenuhi"}}
    lolos = {"checklist": {"lolos_inti": True, "hasil": "ATURAN INTI TERPENUHI"}}
    kep, alasan = lp.keputusan(r, beli, f, nilai=_nilai("AKUMULASI KUAT"), hk_=gagal)
    assert kep == "TUNGGU" and any("Hengky" in a for a in alasan)
    assert lp.keputusan(r, beli, f, nilai=_nilai("AKUMULASI KUAT"), hk_=lolos)[0] == "BELI"
