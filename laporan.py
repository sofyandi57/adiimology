"""Ekspor dokumen riset (PDF) dan data (Excel) dari hasil Adimology + Rekomendasi + Fundamental.

Tanpa Streamlit, sehingga bisa diuji. Edukasi, bukan ajakan transaksi.
"""
from __future__ import annotations

import io

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import pandas as pd  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from reportlab.lib import colors  # noqa: E402
from reportlab.lib.enums import TA_CENTER  # noqa: E402
from reportlab.lib.pagesizes import A4  # noqa: E402
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet  # noqa: E402
from reportlab.lib.units import cm  # noqa: E402
from reportlab.platypus import (CondPageBreak, Image, KeepTogether, Paragraph, SimpleDocTemplate,  # noqa: E402
                                Spacer, Table, TableStyle)

import aliran as al  # noqa: E402
import dinamika as dn  # noqa: E402
import bandar as bd  # noqa: E402
import fundamental as fd  # noqa: E402
import hengky as hk  # noqa: E402
import intraday as idt  # noqa: E402

DISCLAIMER = ("Dokumen ini dibuat otomatis oleh Adimology Calculator dari data Invezgo untuk edukasi dan riset pribadi, "
              "BUKAN ajakan membeli atau menjual efek. Metode Adimology dan aturan teknikal bersifat heuristik dan belum "
              "diuji statistik (backtest). Data laporan keuangan bergantung pada sumber. Keputusan dan risiko transaksi "
              "sepenuhnya tanggung jawab pembaca.")

WARNA_KEP = {"BELI": "#16a34a", "TUNGGU": "#d97706", "HINDARI / JUAL": "#dc2626"}
_FILL = {fd.HIJAU: "DCFCE7", fd.KUNING: "FEF9C3"}


# -------------------------------------------------------------------- keputusan
def _fatal(f: dict | None) -> list[str]:
    if not f:
        return []
    kunci = ("LEBIH BESAR dari total aset", "Ekuitas negatif", "Rugi bersih")
    return [p for p in f.get("peringatan", []) if any(k in p for k in kunci)]


def keputusan(r, rk: dict | None, f: dict | None, alr: dict | None = None,
              nilai: dict | None = None, hk_: dict | None = None, ctx: dict | None = None) -> tuple[str, list[str]]:
    """Keputusan dasar + batas dari konteks (ctx = {'mode': 'cepat'|'lengkap', 'konteks': hasil konteks.ambil}): satu sumber untuk UI, PDF, Excel, TXT, MD."""
    kep, alasan = _keputusan_dasar(r, rk, f, alr, nilai, hk_)
    if ctx:
        import konteks as ktx
        kep, alasan = ktx.terapkan(kep, alasan, ctx.get("konteks"), ctx.get("mode", "lengkap"))
    return kep, alasan


def _keputusan_dasar(r, rk: dict | None, f: dict | None, alr: dict | None = None,
                     nilai: dict | None = None, hk_: dict | None = None) -> tuple[str, list[str]]:
    """BELI hanya bila bukti bandarmologi searah (AKUMULASI KUAT) DAN teknikal di area beli; jebakan atau
    fundamental fatal = HINDARI / JUAL. Target Adimology hanya skenario (belum terbukti backtest)."""
    nilai = nilai or bd.penilaian(r, alr, rk)
    alasan = [f"Bandarmologi: {nilai['kategori']}" + (f" (tingkat bukti: {nilai['tingkat']})" if nilai.get("tingkat") else
                                                       f" (skor {nilai['skor']:+d} dari {nilai['skor_maks']} kriteria terukur)")]
    alasan += [f"Jebakan: {j}" for j in nilai["jebakan"][:2]]
    if rk:
        alasan.append(f"Teknikal: {rk['verdict']} ({rk['alasan']})")
    else:
        alasan.append("Teknikal: tidak terukur (data gagal diambil)")
    fatal = _fatal(f)
    if f is None:
        alasan.append("Fundamental: tidak terukur (data gagal diambil)")
    elif fatal:
        alasan.append("Fundamental: " + fatal[0])
    else:
        alasan.append(f"Fundamental: tidak ada tanda bahaya besar ({len(f.get('peringatan', []))} catatan kuning)")
    alasan.append(f"Skenario Adimology: {r.verdict} (bukan dasar keputusan; belum terbukti di backtest)")

    if fatal or bd.gate(nilai) == "waspada" or (rk and rk["verdict"] == "HINDARI"):
        return "HINDARI / JUAL", alasan
    if hk_ is not None:
        alasan.append(f"Aturan Hengky: {hk_['checklist']['hasil']}")
    if bd.gate(nilai) == "kuat" and rk and rk["verdict"] == "BELI DI AREA ENTRY":
        if hk_ is not None and not hk_["checklist"]["lolos_inti"]:
            alasan.append("BELI ditahan: aturan inti Hengky belum terpenuhi (uptrend, likuid, level, level pembatal)")
            return "TUNGGU", alasan
        return "BELI", alasan
    if rk and rk["verdict"] == "BELI DI AREA ENTRY":
        alasan.append("BELI ditahan: bukti bandarmologi belum cukup (perlu AKUMULASI TERKONFIRMASI oleh perilaku harga)")
    return "TUNGGU", alasan


# --------------------------------------------------------------------- helpers
def _n(v, d=0):
    return "-" if v is None or pd.isna(v) else f"{v:,.{d}f}"


def _pct(v):
    return "-" if v is None or pd.isna(v) else f"{v:+.1%}"


def _safe(t) -> str:
    """Font standar PDF (Helvetica) tidak punya beberapa simbol."""
    rep = {"−": "-", "→": "->", "≥": ">=", "≤": "<=", "±": "+/-", "–": "-", "—": "-", "×": "x", "≠": "!="}
    s = str(t)
    for a, b in rep.items():
        s = s.replace(a, b)
    return s.encode("latin-1", "ignore").decode("latin-1").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _area(rk: dict) -> str:
    return f"{_n(rk['entry_lo'])} - {_n(rk['entry_hi'])}" if rk["zona_valid"] else "tidak ada (R/R kurang)"


def ringkasan_rows(r, rk: dict | None, f: dict | None, waktu: str, alr: dict | None = None,
                   hk_: dict | None = None, dinamis: dict | None = None, ctx: dict | None = None) -> list[tuple[str, str]]:
    kep, _ = keputusan(r, rk, f, alr, bd.penilaian(r, alr, rk, dinamis), hk_, ctx)
    rows = [("Kode", r.params.kode), ("Nama", (f or {}).get("nama") or "-"), ("Waktu analisa", waktu),
            ("Keputusan", kep), ("Harga terakhir", _n(r.harga)),
            ("Penilaian bandarmologi", bd.penilaian(r, alr, rk, dinamis)["kategori"]),
            ("Skenario Adimology (belum terbukti)", r.verdict), ("Bandar (broker)", f"{r.bandar['code']} ({r.bandar.get('name', '')})"),
            ("Avg beli bandar", _n(r.bandar["avg"], 1)), ("Net bandar (lot)", _n(r.bandar["net_lot"])),
            ("Target realistis (Adimology)", _n(r.calc["t_real_r"])), ("Target max (Adimology)", _n(r.calc["t_max_r"])),
            ("SL low window (Adimology)", _n(r.plan["sl"])),
            ("Window broker summary", f"{r.tgl_awal} s/d {r.tgl_akhir} ({r.params.window_hari} hari)")]
    if rk:
        rows += [("Verdict teknikal", rk["verdict"]), ("Area beli", _area(rk)),
                 (f"Entry maksimal (R/R {rk['min_rr']:.1f})", _n(rk["entry_maks"])),
                 ("Stop loss (teknikal)", _n(rk["sl"])), ("Area jual (TP1 - TP2)", f"{_n(rk['tp1'])} - {_n(rk['tp2'])}"),
                 ("Target utama", f"{rk['target_nama']} {_n(rk['target'])}"),
                 ("R/R rata-rata area beli", _n(rk["rr"], 2)), ("RSI 14", _n(rk["rsi"], 1)),
                 ("Bandar value (status top-5)", rk["bandar"].get("status", "-"))]
    if f:
        rows.append(("Laporan keuangan terakhir", f["periode"]))
    return rows


# ------------------------------------------------------------------------ Excel
def excel_bytes(r, rk: dict | None, f: dict | None, waktu: str, ovn: dict | None = None, alr: dict | None = None,
                b200: dict | None = None, hk_: dict | None = None, ritel: dict | None = None, dinamis: dict | None = None,
                ctx: dict | None = None) -> bytes:
    buf = io.BytesIO()
    nilai_ = bd.penilaian(r, alr, rk, dinamis)
    kep, alasan = keputusan(r, rk, f, alr, nilai_, hk_, ctx)
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        ring = pd.DataFrame(ringkasan_rows(r, rk, f, waktu, alr, hk_, dinamis, ctx), columns=["Item", "Nilai"])
        ring = pd.concat([ring, pd.DataFrame([("Alasan keputusan", a) for a in alasan], columns=["Item", "Nilai"]),
                          pd.DataFrame([("Disclaimer", DISCLAIMER)], columns=["Item", "Nilai"])])
        ring.to_excel(xw, sheet_name="Ringkasan", index=False)
        if rk:
            ind = rk["ind"].tail(250)
            cols = ["tgl", "open", "high", "low", "close", "volume", "ma5", "ma20", "ma50", "ma100", "ma200",
                    "rsi", "macd", "signal", "hist", "atr"]
            ind[[c for c in cols if c in ind]].to_excel(xw, sheet_name="Teknikal", index=False)
            sr = [{"Jenis": "Resistance", "Harga": c["harga"], "Sentuhan": c["n"]} for c in rk["res"]] + \
                 [{"Jenis": "Support", "Harga": c["harga"], "Sentuhan": c["n"]} for c in rk["sup"]]
            pd.DataFrame(sr or [{"Jenis": "-", "Harga": None, "Sentuhan": None}]).to_excel(xw, sheet_name="Support Resistance", index=False)
            pd.DataFrame([{"Poin": p, "Sinyal": t} for p, t in rk["poin"]]).to_excel(xw, sheet_name="Sinyal", index=False)
        bcols = [c for c in ["code", "name", "net_lot", "porsi_net_%", "buy_avg", "sell_avg", "net_value"] if c in r.bs]
        r.bs.sort_values("net_volume", ascending=False)[bcols].to_excel(xw, sheet_name="Broker", index=False)
        r.harian.to_excel(xw, sheet_name="Bandar harian", index=False)
        nilai = bd.penilaian(r, alr, rk, dinamis)
        pd.concat([bd.tabel(nilai), pd.DataFrame([{"Kriteria": "JEBAKAN", "Status": "⚠", "Bukti": j} for j in nilai["jebakan"]]),
                   pd.DataFrame([{"Kriteria": "Kategori", "Status": nilai["kategori"], "Bukti": (f"tingkat bukti: {nilai['tingkat']}" if nilai.get("tingkat") else f"skor {nilai['skor']:+d}")}])]
                  ).to_excel(xw, sheet_name="Penilaian bandar", index=False)
        if b200:
            pd.DataFrame([dict(b200, jenis="ringkasan", akumulator=None, distributor=None, konsentrasi=b200["konsentrasi"]["status"])]
                         ).to_excel(xw, sheet_name="Bandar jangka panjang", index=False)
            pd.concat([pd.DataFrame(b200["akumulator"]).assign(peran="akumulator"),
                       pd.DataFrame(b200["distributor"]).assign(peran="distributor")]
                      ).to_excel(xw, sheet_name="Bandar panjang broker", index=False)
        if hk_:
            hk.tabel(hk_["checklist"]).to_excel(xw, sheet_name="Aturan Hengky", index=False)
            if hk_["pv"]:
                hk_["pv"]["harian"].assign(label=hk_["pv"]["harian"]["label"].map(hk.PV_TEKS)).to_excel(
                    xw, sheet_name="Harga-volume 5h", index=False)
        if ritel:
            pd.DataFrame([{"status": ritel["status"], "minat": ritel["minat"], "ringkas": ritel["ringkas"]}] +
                         [{"status": "", "minat": "", "ringkas": x} for x in ritel["poin"]]).to_excel(xw, sheet_name="Minat ritel", index=False)
        if dinamis and dinamis.get("dimensi"):
            pd.DataFrame([dict(Dimensi=x["dimensi"], Bobot=x["bobot"], Skor=round(x["skor"]), Mendukung="; ".join(x["mendukung"]), Menentang="; ".join(x["menentang"]))
                          for x in dinamis["dimensi"]]).to_excel(xw, sheet_name="Bandar dinamis", index=False)
            dinamis["tabel"].to_excel(xw, sheet_name="Bandar per jendela", index=False)
        pd.DataFrame(al.ringkas(r, alr)).to_excel(xw, sheet_name="Konfirmasi aliran", index=False)
        pd.DataFrame([{"Catatan": bd.HASIL_BACKTEST}]).to_excel(xw, sheet_name="Hasil backtest", index=False)
        if alr is not None and not alr["konsistensi"].empty:
            alr["konsistensi"].to_excel(xw, sheet_name="Konsistensi bandar", index=False)
        if ovn:
            idt.tabel(ovn["bsjp"], bsjp=True).to_excel(xw, sheet_name="BSJP", index_label="Metrik")
            idt.tabel(ovn["bpjp"], bsjp=False).to_excel(xw, sheet_name="BPJP", index_label="Metrik")
            ovn["data_bsjp"].to_excel(xw, sheet_name="BSJP data harian", index=False)
            ovn["data_bpjp"].to_excel(xw, sheet_name="BPJP data harian", index=False)
        if f:
            for nama, kunci in [("PnL", "pnl"), ("Neraca", "bs"), ("Arus kas", "cf")]:
                f[kunci].to_excel(xw, sheet_name=nama, index_label="Rp miliar / rasio")
            f["val"].to_excel(xw, sheet_name="Valuasi", index=False)
            pd.DataFrame({"Peringatan (kuning)": pd.Series(f["peringatan"], dtype=object),
                          "Sorotan (hijau)": pd.Series(f["sorotan"], dtype=object)}).to_excel(xw, sheet_name="Catatan fundamental", index=False)
        wb = xw.book
        for ws in wb.worksheets:
            for c in ws[1]:
                c.font = Font(bold=True)
            for i, col in enumerate(ws.columns, 1):
                w = max(len(str(c.value)) if c.value is not None else 0 for c in col)
                ws.column_dimensions[get_column_letter(i)].width = min(max(10, w + 2), 60)
        ws = wb["Ringkasan"]
        ws.column_dimensions["B"].width = 90
        for row in ws.iter_rows(min_row=2):
            row[1].alignment = Alignment(wrap_text=True, vertical="top")
            if row[0].value == "Keputusan":
                row[1].font = Font(bold=True, color=WARNA_KEP[kep].lstrip("#"))
        if f:
            for nama, kunci in [("PnL", "pnl"), ("Neraca", "bs"), ("Arus kas", "cf")]:
                ws = wb[nama]
                idx = {str(ws.cell(row=i, column=1).value): i for i in range(2, ws.max_row + 1)}
                cols = {str(ws.cell(row=1, column=j).value): j for j in range(2, ws.max_column + 1)}
                for (t, rn, cn), css in f["mark"].items():
                    if t == kunci and rn in idx and cn in cols:
                        ws.cell(row=idx[rn], column=cols[cn]).fill = PatternFill("solid", fgColor=_FILL[css])
    return buf.getvalue()


# -------------------------------------------------------------------------- PDF
def _chart_png(r, rk: dict | None) -> bytes:
    fig, ax = plt.subplots(figsize=(10, 3.6), dpi=130)
    if rk:
        d = rk["ind"].tail(160)
        x = range(len(d))
        ax.plot(x, d["close"], color="#111827", lw=1.4, label="Close")
        for n, c in {20: "#f59e0b", 50: "#10b981", 200: "#ef4444"}.items():
            if d[f"ma{n}"].notna().any():
                ax.plot(x, d[f"ma{n}"], color=c, lw=1, label=f"MA{n}")
        if rk["zona_valid"]:
            ax.axhspan(rk["entry_lo"], rk["entry_hi"], color="#2563eb", alpha=0.15, label="Area beli")
        ax.axhspan(rk["tp1"], rk["tp2"], color="#16a34a", alpha=0.10, label="Area jual")
        lines = [(rk["tp1"], "TP1", "#16a34a"), (rk["tp2"], "TP2", "#15803d"), (rk["sl"], "SL", "#dc2626")]
        ticks = d["tgl"].astype(str).tolist()
    else:
        d = r.chart.tail(60)
        x = range(len(d))
        ax.plot(x, d["close"], color="#111827", lw=1.4, label="Close")
        lines = []
        ticks = d["tgl"].astype(str).tolist()
    lines += [(r.calc["t_real_r"], "Target realistis", "#d97706"), (r.bandar["avg"], "Avg bandar", "#2563eb")]
    for y, nama, c in lines:
        ax.axhline(y, color=c, lw=0.9, ls="--")
        ax.text(len(d) - 1, y, f" {nama} {y:,.0f}", color=c, fontsize=7, va="bottom", ha="right")
    step = max(1, len(ticks) // 8)
    ax.set_xticks(list(x)[::step])
    ax.set_xticklabels(ticks[::step], fontsize=7, rotation=0)
    ax.tick_params(axis="y", labelsize=7)
    ax.grid(alpha=0.25)
    ax.legend(fontsize=7, loc="upper left", ncol=5)
    ax.set_title(f"{r.params.kode}: harga, MA, area entry, target & SL", fontsize=9)
    fig.tight_layout()
    out = io.BytesIO()
    fig.savefig(out, format="png")
    plt.close(fig)
    return out.getvalue()


def _tbl(data, widths=None, header=True, cell_bg: dict | None = None, font=8):
    t = Table(data, colWidths=widths, repeatRows=1 if header else 0)
    st = [("FONT", (0, 0), (-1, -1), "Helvetica", font), ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#d1d5db")),
          ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]
    if header:
        st += [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2937")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
               ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", font)]
    for (ri, ci), c in (cell_bg or {}).items():
        st.append(("BACKGROUND", (ci, ri), (ci, ri), colors.HexColor(c)))
    t.setStyle(TableStyle(st))
    return t


def _fund_table(df: pd.DataFrame, mark: dict, kunci: str, lebar: float):
    hdr = [""] + list(df.columns)
    data = [hdr] + [[_safe(i)] + [_safe(v) for v in row] for i, row in zip(df.index, df.values)]
    bg = {}
    for (t, rn, cn), css in mark.items():
        if t == kunci and rn in df.index and cn in df.columns:
            bg[(list(df.index).index(rn) + 1, list(df.columns).index(cn) + 1)] = "#" + _FILL[css]
    first = 4.2 * cm
    rest = (lebar - first) / max(1, len(df.columns))
    return _tbl(data, [first] + [rest] * len(df.columns), cell_bg=bg, font=7)


def pdf_bytes(r, rk: dict | None, f: dict | None, waktu: str, ovn: dict | None = None, alr: dict | None = None,
              b200: dict | None = None, hk_: dict | None = None, ritel: dict | None = None, dinamis: dict | None = None,
              ctx: dict | None = None) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=1.6 * cm, rightMargin=1.6 * cm, topMargin=1.4 * cm,
                            bottomMargin=1.4 * cm, title=f"Riset {r.params.kode}", author="Adimology Calculator")
    W = A4[0] - 3.2 * cm
    ss = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=ss["Title"], fontSize=18, spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontSize=12, spaceBefore=10, spaceAfter=4)
    p = ParagraphStyle("p", parent=ss["BodyText"], fontSize=8.5, leading=11)
    small = ParagraphStyle("s", parent=p, fontSize=7, textColor=colors.HexColor("#6b7280"))
    nilai_ = bd.penilaian(r, alr, rk, dinamis)
    kep, alasan = keputusan(r, rk, f, alr, nilai_, hk_, ctx)
    big = ParagraphStyle("k", parent=ss["Title"], fontSize=22, textColor=colors.white, alignment=TA_CENTER, leading=26)

    el = [Paragraph(_safe(f"Riset Saham {r.params.kode}" + (f" - {f['nama']}" if f and f.get("nama") else "")), h1),
          Paragraph(_safe(f"Dibuat {waktu} · Adimology Calculator · data Invezgo"), small), Spacer(1, 8)]
    box = Table([[Paragraph(f"KEPUTUSAN: {_safe(kep)}", big)]], colWidths=[W])
    box.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(WARNA_KEP[kep])),
                             ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 10)]))
    el += [box, Spacer(1, 6)]
    el += [Paragraph("<b>Dasar keputusan</b>", p)] + [Paragraph("- " + _safe(a), p) for a in alasan]

    el.append(Paragraph("Rencana trading", h2))
    if rk:
        plan = [["Harga", "Area beli", f"Entry maks (R/R {rk['min_rr']:.1f})", "Stop loss", "Area jual (TP1-TP2)", "R/R"],
                [_n(rk["harga"]), _area(rk), _n(rk["entry_maks"]), _n(rk["sl"]),
                 f"{_n(rk['tp1'])} - {_n(rk['tp2'])} ({_pct(rk['tp1'] / rk['harga'] - 1)} s/d {_pct(rk['tp2'] / rk['harga'] - 1)})",
                 _n(rk["rr"], 2)]]
        el.append(_tbl(plan, [W * 0.11, W * 0.16, W * 0.16, W * 0.11, W * 0.34, W * 0.12]))
        el.append(Paragraph(_safe(f"Area beli = {rk['entry_basis']} s/d entry maksimal yang masih memberi R/R {rk['min_rr']:.1f} ke "
                                  f"{rk['target_nama']} {_n(rk['target'])}; SL = support - 1 ATR14; area jual dari {rk['tp_basis']}."), small))
        for ev in rk.get("split", []):
            el.append(Paragraph(_safe(f"Harga sebelum {ev['tgl']} disesuaikan untuk {ev['jenis']}."), small))
    else:
        el.append(Paragraph("Analisa teknikal tidak tersedia.", p))
    el += [Spacer(1, 6), Image(io.BytesIO(_chart_png(r, rk)), width=W, height=W * 0.36)]

    nilai = bd.penilaian(r, alr, rk, dinamis)
    el.append(Paragraph(_safe(f"Bandarmologi: {nilai['kategori']}" + (f" - tingkat bukti: {nilai['tingkat']}" if nilai.get("tingkat") else
                                                                       f" (skor {nilai['skor']:+d} dari {nilai['skor_maks']} kriteria terukur)")), h2))
    if dinamis and dinamis.get("dimensi"):
        el.append(_tbl([["Dimensi (bobot = usulan desain)", "Skor", "Mendukung", "Menentang"]] + [
            [Paragraph(_safe(f"{x['dimensi']} ({x['bobot']:.0%})"), p), f"{x['skor']:.0f}", Paragraph(_safe("; ".join(x["mendukung"])[:160]), p),
             Paragraph(_safe("; ".join(x["menentang"])[:160]), p)] for x in dinamis["dimensi"]], [W * 0.28, W * 0.08, W * 0.32, W * 0.32]))
        el.append(Paragraph(_safe(f"Keyakinan {dinamis['keyakinan']}. Pembatalan: " + "; ".join(dinamis["pembatalan"]) + ". " + dn.CATATAN_BOBOT), small))
    ikon = {bd.YA: "ya", bd.TIDAK: "TIDAK", bd.NETRAL: "netral", bd.NA: "tidak diukur"}
    bg = {(i + 1, 1): {bd.YA: "#DCFCE7", bd.TIDAK: "#FEE2E2", bd.NETRAL: "#F3F4F6", bd.NA: "#F3F4F6"}[x["status"]]
          for i, x in enumerate(nilai["kriteria"])}
    el.append(_tbl([["Kriteria", "Status", "Bukti"]] +
                   [[Paragraph(_safe(x["kriteria"] + (" *" if x["inti"] else "")), p), ikon[x["status"]], Paragraph(_safe(x["bukti"]), p)]
                    for x in nilai["kriteria"]], [W * 0.3, W * 0.12, W * 0.58], cell_bg=bg))
    for j in nilai["jebakan"]:
        el.append(_tbl([[Paragraph("JEBAKAN: " + _safe(j), p)]], [W], header=False, cell_bg={(0, 0): "#FEE2E2"}))
    if ritel:
        el.append(_tbl([[Paragraph(_safe(f"Minat ritel: {ritel['status']} (minat {ritel['minat']}). {ritel['ringkas']} "
                                         + " ".join(ritel["poin"])), p)]], [W], header=False, cell_bg={(0, 0): "#E0F2FE"}))
    for i_ in nilai["info"]:
        el.append(Paragraph(_safe("Info: " + i_), small))
    el.append(Paragraph(_safe("* kriteria inti. Hasil backtest Adimology: " + bd.HASIL_BACKTEST), small))
    if b200:
        el.append(_tbl([[Paragraph(_safe(
            f"Bandar {b200['hari']} hari bursa ({b200['dari']} s/d {b200['sampai']}): {b200['status']} - {b200['nyangkut']}. "
            f"{b200['keterangan']}."), p)]], [W], header=False,
            cell_bg={(0, 0): {"NYANGKUT": "#FEF9C3", "TIDAK NYANGKUT": "#DCFCE7"}.get(b200["nyangkut"], "#F3F4F6")}))
        data = [["Akumulator", "Net lot", "Avg beli", "Harga vs avg"]] + [
            [_safe(x["code"]), _n(x["net_lot"]), _n(x["avg"], 1), _pct(x["floating"])] for x in b200["akumulator"]]
        el.append(_tbl(data, [W * 0.25] * 4))
    el.append(Paragraph("Skenario target Adimology (belum terbukti backtest)", h2))
    ad_rows = [["Verdict", "Bandar", "Avg bandar", "Net (lot)", "Target realistis", "Target max", "SL window"],
               [Paragraph(_safe(r.verdict), ParagraphStyle("c", parent=p, fontSize=8, leading=9)), _safe(r.bandar["code"]), _n(r.bandar["avg"], 1), _n(r.bandar["net_lot"]),
                f"{_n(r.calc['t_real_r'])} ({_pct(r.calc['upside_real'])})", f"{_n(r.calc['t_max_r'])} ({_pct(r.calc['upside_max'])})",
                _n(r.plan["sl"])]]
    el.append(_tbl(ad_rows, [W * 0.19, W * 0.09, W * 0.12, W * 0.12, W * 0.18, W * 0.18, W * 0.12]))
    el += [Paragraph("- " + _safe(a), p) for a in r.alasan]
    el.append(Paragraph(_safe(f"Rentang target {_n(r.calc['t_real_r'])} - {_n(r.calc['t_max_r'])}; lot per papan "
                              f"{_n(r.lot_per_papan)} ({r.lpp_sumber}); mode bandar {r.params.mode_bandar}."), small))
    top = r.bs.sort_values("net_volume", ascending=False)
    bt = [["Top net buy", "Net lot", "Avg beli", "Top net sell", "Net lot", "Avg jual"]]
    buy, sell = top.head(5).reset_index(drop=True), top.tail(5).iloc[::-1].reset_index(drop=True)
    for i in range(max(len(buy), len(sell))):
        b = buy.iloc[i] if i < len(buy) else None
        s_ = sell.iloc[i] if i < len(sell) else None
        bt.append([_safe(b["code"]) if b is not None else "", _n(b["net_lot"]) if b is not None else "",
                   _n(b["buy_avg"], 1) if b is not None else "", _safe(s_["code"]) if s_ is not None else "",
                   _n(s_["net_lot"]) if s_ is not None else "", _n(s_["sell_avg"], 1) if s_ is not None else ""])
    el += [Spacer(1, 4), _tbl(bt, [W / 6] * 6)]

    if rk:
        el.append(Paragraph("Teknikal", h2))
        ma = [["MA5", "MA20", "MA50", "MA100", "MA200", "RSI 14", "MACD / signal"],
              [_n(rk["ma"][5]), _n(rk["ma"][20]), _n(rk["ma"][50]), _n(rk["ma"][100]), _n(rk["ma"][200]),
               _n(rk["rsi"], 1), f"{_n(rk['macd'], 2)} / {_n(rk['signal'], 2)}"]]
        el.append(_tbl(ma, [W / 7] * 7))
        sr = [["Support", "Resistance"]] + [[_n(a["harga"]) if a else "", _n(b["harga"]) if b else ""]
                                            for a, b in zip((rk["sup"] + [None] * 3)[:3], (rk["res"] + [None] * 3)[:3])]
        bv = rk["bandar"]
        bvt = [["Bandar value", "Rp miliar"], ["Net bandar terpilih", _n((bv.get("bandar") or 0) / 1e9, 1)],
               ["Net top-5 akum. vs dist.", _n((bv.get("net_top5") or 0) / 1e9, 1)], ["Status", _safe(bv.get("status", "-"))]]
        pair = Table([[_tbl(sr, [W * 0.22] * 2), _tbl(bvt, [W * 0.28, W * 0.18])]], colWidths=[W * 0.48, W * 0.52])
        el += [Spacer(1, 4), pair, Spacer(1, 4)]
        sig = [Paragraph(_safe(("[+] " if s > 0 else "[-] " if s < 0 else "[ ] ") + t), p) for s, t in rk["poin"]]
        half = (len(sig) + 1) // 2
        kiri, kanan = sig[:half], sig[half:] + [""] * (half - len(sig[half:]))
        sigt = Table([[a, b] for a, b in zip(kiri, kanan)], colWidths=[W / 2] * 2)
        sigt.setStyle(TableStyle([("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
                                  ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
        el.append(KeepTogether([sigt, Paragraph(_safe(f"Skor {rk['skor']:+d} dari +/-{rk['skor_max']}. "
                                                      f"Verdict teknikal: {rk['verdict']} - {rk['alasan']}."), p)]))

    if hk_:
        cl = hk_["checklist"]
        el += [CondPageBreak(7 * cm), Paragraph(_safe("Aturan disiplin trading (materi Hengky / Remora)"), h2),
               Paragraph(_safe(f"{cl['hasil']} · profil: {hk_['profil']['profil']} ({hk_['profil']['alasan']})"), p)]
        bg = {(i + 1, 1): {hk.YA: "#DCFCE7", hk.TIDAK: "#FEE2E2", hk.NETRAL: "#F3F4F6", hk.NA: "#F3F4F6"}[k["status"]]
              for i, k in enumerate(cl["kriteria"])}
        ik = {hk.YA: "ya", hk.TIDAK: "TIDAK", hk.NETRAL: "netral", hk.NA: "tidak diukur"}
        el.append(_tbl([["Aturan", "Status", "Bukti"]] + [
            [Paragraph(_safe(k["kriteria"] + (" *" if k["inti"] else "")), p), ik[k["status"]], Paragraph(_safe(k["bukti"]), p)]
            for k in cl["kriteria"]], [W * 0.32, W * 0.1, W * 0.58], cell_bg=bg))
        rn = hk_["rencana"]
        if rn and (rn["support"] or rn["breakout"]):
            rows = [["Rencana", "Level", "Entry", "Stop 1 tick", "Target", "R/R tick | ATR"]]
            if rn["support"]:
                s_ = rn["support"]
                rows.append(["Beli di support", _n(s_["level"]), _n(s_["entry_ref"]), _n(s_["stop_tick"]), _n(s_["target"]),
                             f"{s_['rr']:.1f} | {s_['rr_atr']:.1f}"])
            if rn["breakout"]:
                b_ = rn["breakout"]
                rows.append(["Breakout", _n(b_["level"]), _n(b_["entry"]), _n(b_["stop_tick"]), _n(b_["target"]),
                             f"{b_['rr']:.1f} | {b_['rr_atr']:.1f}"])
            el += [Spacer(1, 4), _tbl(rows, [W * 0.2, W * 0.15, W * 0.15, W * 0.17, W * 0.15, W * 0.18])]
            el.append(Paragraph(_safe(f"Harga impas setelah fee {_n(rn['impas'])} ({rn['impas_pct']:+.2%}). Stop 1 tick sangat ketat; "
                                      "jangan average down. " + hk.CATATAN_KRITIS), small))
    if ovn:
        el += [CondPageBreak(7 * cm), Paragraph(_safe(f"Statistik BPJP / BSJP ({ovn['dari']} s/d {ovn['sampai']}, candle 30 menit)"), h2)]
        sg = ovn["sinyal"].get("bsjp")
        if sg:
            warna = "#DCFCE7" if sg["layak"] else "#FEF9C3"
            el.append(_tbl([[Paragraph(_safe(
                f"BSJP {ovn['hari_ini']['tgl']}: volume jam terakhir {sg['rvol']:.1f}x median ({sg['kelompok']}). "
                f"Peluang gap naik {sg['p_gap']:.0%}, sentuh +1% bersih fee {sg['p_target']:.0%}. Target jual {sg['target']:,.0f} "
                f"({sg['naik']:+.1%}); puncak paling sering {idt.label_jam(sg['jam_jual'])} ({sg['p_jam']:.0%})."
                + ("" if sg["layak"] else " Peluang historis kurang: BSJP tidak disarankan.")), p)]], [W], header=False,
                cell_bg={(0, 0): warna}))
        el.append(Paragraph(_safe(f"Ambang volume tinggi: sore (15:00-16:00) >= {ovn['ambang']['sore']:,.0f} lembar; "
                                  f"30 menit pertama >= {ovn['ambang']['buka']:,.0f} lembar. Fee {ovn['fee']:.2%} sudah diperhitungkan."), small))
        for judul, blok, b in (("BSJP: beli close, jual sesi 1 besok", ovn["bsjp"], True),
                               ("BPJP: beli ~09:30, jual s/d akhir sesi 1", ovn["bpjp"], False)):
            t = idt.tabel(blok, bsjp=b)
            data = [["Metrik"] + list(t.columns)] + [[_safe(i)] + [_safe(v) for v in row] for i, row in zip(t.index, t.values)]
            el.append(KeepTogether([Paragraph(f"<b>{_safe(judul)}</b>", p),
                                    _tbl(data, [W * 0.4] + [W * 0.2] * len(t.columns), font=7), Spacer(1, 4)]))
    if f:
        el += [CondPageBreak(6 * cm), Paragraph(_safe(f"Fundamental ({f['periode']})"), h2),
               Paragraph(_safe(" · ".join(x for x in [f["sektor"], f["subsektor"]] if x) +
                               " · nominal Rp miliar per kuartal · hijau = tumbuh, kuning = perlu dicek"), small)]
        for w_ in f["peringatan"]:
            el.append(_tbl([[Paragraph("PERHATIAN: " + _safe(w_), p)]], [W], header=False, cell_bg={(0, 0): "#FEF9C3"}))
        for s_ in f["sorotan"]:
            el.append(_tbl([[Paragraph("TUMBUH: " + _safe(s_), p)]], [W], header=False, cell_bg={(0, 0): "#DCFCE7"}))
        for judul, kunci in [("Laba rugi", "pnl"), ("Neraca", "bs"), ("Arus kas", "cf")]:
            el.append(KeepTogether([Paragraph(f"<b>{judul}</b>", p), _fund_table(f[kunci], f["mark"], kunci, W), Spacer(1, 6)]))
        vals = list(f["val"].values)
        half = (len(vals) + 1) // 2
        sub = []
        for part in (vals[:half], vals[half:]):
            bg = {(i + 1, 1): "#" + _FILL[f["vmark"][m]] for i, (m, _) in enumerate(part) if m in f["vmark"]}
            sub.append(_tbl([["Metrik", "Nilai"]] + [[_safe(a), _safe(b)] for a, b in part], [W * 0.25, W * 0.22], cell_bg=bg))
        el.append(KeepTogether([Paragraph("<b>Valuasi</b>", p), Table([sub], colWidths=[W / 2] * 2)]))

    el.append(KeepTogether([Spacer(1, 8), Paragraph("<b>Disclaimer</b>", p), Paragraph(_safe(DISCLAIMER), small)]))

    def footer(canvas, d):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#6b7280"))
        canvas.drawString(1.6 * cm, 0.8 * cm, f"Riset {r.params.kode} · {waktu} · edukasi, bukan ajakan transaksi")
        canvas.drawRightString(A4[0] - 1.6 * cm, 0.8 * cm, f"Hal. {d.page}")
        canvas.restoreState()

    doc.build(el, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()
